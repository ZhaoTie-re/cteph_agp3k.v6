#!/usr/bin/env python3
# ---------------------------------------------------------------------------
# Purpose : The snapshot workbook of the CTEPH knowledge graph, organised by
#           question. Each question sheet opens with its answer in brief, then
#           a nested table: group > record > source, opened with the outline
#           buttons (+ / -, levels 1 2 3). The appendix holds the flat tables,
#           with filters. House style of the integration tables
#           (integration/gene_signals): booktabs rules, italic gene symbols, P as
#           m x 10^-e, lettered footnotes, the reason for every "—", one
#           Abbreviations line. Presentation only: every value is taken as is
#           from the tables; generated summaries are labelled.
#             Guide                  the questions, how to read, the snapshot
#             1 Related conditions   category > condition > evidence (quotes,
#                                    ontology, shared GWAS signals)
#             2 Shared genes         gene x condition map of genetic evidence
#             3 CTEPH genetics       the CTEPH GWAS loci; the literature by gene
#                                    theme > gene > finding
#             4 Study methods        design > study, its methods by step
#             5 Condition genetics   category > condition > gene
#             6 Mechanisms           Reactome theme > target > publication
#             7 East Asian           kind of evidence > item
#             8 Changes              since the previous snapshot
#             A1-A6                  GWAS studies, GWAS loci, colocalisation,
#                                    curated genes, CTEPH studies, screening
#           The workbook names no file, script, parameter or pipeline.
# Inputs  : the tidy tables, summaries, changes, knowledge JSON, gene themes
# Outputs : the workbook, workbook.json (layout and rows per sheet)
# ---------------------------------------------------------------------------
import argparse
import json
import math
import re
import sys
import warnings
import zipfile
from pathlib import Path

import pandas as pd
import xlsxwriter

sys.path.insert(0, str(Path(__file__).resolve().parent))
from ot_common import (CURATED, MISSING, NO_GENE, OTHER_GENES, ancestry_label, assign_themes,  # noqa: E402
                       gene_tallies, is_east_asian, is_japanese, load_themes, log, num, pub_url, read_tsv,
                       study_population, tally_text, theme_order, variant_label)

FORBIDDEN_TOKENS = ['.tsv', '.py', '.nf', '.yaml', '/LARGE', '/home', 'nextflow', 'europepmc', 'clinical_precedence',
                    'gwas_credible_sets', '/snapshots', 'FetchDate', 'ensembl_id', 'l2g1', 'datasource', '_mainland',
                    'node_id', 'custom:']
FONT = 'Arial'
INK, INK2, GREY, RULE_GREY, NOTE = '#1A1A1A', '#595959', '#7F7F7F', '#BFBFBF', '#404040'
LINK = '#1F4E79'
GROUP0, GROUP1 = '#DCE6F1', '#F2F2F2'        # rows of outline level 0 / rows that open
KIND = {'GWAS': 'G', 'Mendelian': 'C', 'Literature': 'L'}
KIND_FILL = {'G': '#C6DBEF', 'C': '#C7E9C0', 'L': '#FCBBA1'}
MULTI_FILL = '#8C8C8C'
SOURCE_FILL = {'full text': '#C7E9C0', 'abstract': '#EDEDED'}
THIN, MEDIUM = 1, 2
CORE = 'EFO_0022954'

QUESTIONS = [('1 Related conditions', 'Which conditions are related to CTEPH, and on what evidence?'),
             ('2 Shared genes', 'Which genes do CTEPH and its related conditions share?'),
             ('3 CTEPH genetics', 'What is known about the genetics of CTEPH?'),
             ('4 Study methods', 'How were the genetic studies of CTEPH done?'),
             ('5 Condition genetics', 'What are the genetics of each related condition?'),
             ('6 Mechanisms', 'Which mechanisms does the CTEPH literature describe?'),
             ('7 East Asian', 'What evidence comes from East Asian and Japanese populations?'),
             ('8 Changes', 'What changed since the previous snapshot?')]
APPENDIX = [('A1 GWAS studies', 'Every GWAS study of CTEPH and the related conditions: design, size, ancestry'),
            ('A2 GWAS loci', 'Every credible set: lead variant, P, β, fine-mapping, locus-to-gene genes'),
            ('A3 Colocalisation', 'CTEPH loci × colocalising GWAS signals'),
            ('A4 Curated genes', 'Curated gene evidence of the related conditions, by source'),
            ('A5 CTEPH studies', 'The genetic studies of CTEPH, one row each'),
            ('A6 Screening', 'Every PubMed record screened for the CTEPH literature, and the decision')]
SH = ['Guide'] + [s for s, _q in QUESTIONS] + [s for s, _d in APPENDIX]
CATEGORY_ORDER = ['Precursor / spectrum', 'Risk factor', 'Comorbidity', 'Intermediate trait', 'PH group']
CATEGORY_DESC = {
    'Precursor / spectrum': 'Venous thromboembolism and the chronic thromboembolic disease spectrum CTEPH develops '
                            'from.',
    'Risk factor': 'Conditions and procedures that raise the risk of CTEPH (mostly against non-thromboembolic PH '
                   'controls).',
    'Comorbidity': 'Conditions more frequent in CTEPH, or reported with it.',
    'Intermediate trait': 'Blood traits between genotype and thrombosis.',
    'PH group': 'Pulmonary hypertension and its other forms: the neighbours of CTEPH in the disease ontology.',
    'multi': 'Traits without a clinical or ontology relation whose GWAS signal colocalises with two or more CTEPH '
             'loci.',
    'single': 'Traits whose GWAS signal colocalises with one CTEPH locus only, mostly the pleiotropic ABO locus: '
              'weak evidence of a relation.'}
SHORT = {'chronic thromboembolic pulmonary hypertension': 'CTEPH', 'Venous thromboembolism': 'VTE',
         'Deep vein thrombosis': 'DVT', 'Pulmonary embolism': 'PE', 'Pulmonary arterial hypertension': 'PAH',
         'Pulmonary arterial hypertension associated with another disease': 'PAH, associated',
         'heritable pulmonary arterial hypertension': 'Heritable PAH',
         'idiopathic pulmonary arterial hypertension': 'Idiopathic PAH',
         'pulmonary veno-occlusive disease and/or pulmonary capillary haemangiomatosis': 'PVOD / PCH',
         'hyperuricemia-pulmonary hypertension-renal failure-alkalosis syndrome': 'HUPRA syndrome',
         'Hypothyroidism / thyroid replacement therapy': 'Hypothyroidism',
         'Sickle cell disease and haemoglobinopathies': 'Sickle cell disease',
         'Fibrinogen (level and variants)': 'Fibrinogen', 'Factor XI level (F11 locus)': 'Factor XI level',
         'Antiphospholipid syndrome / antiphospholipid antibodies': 'Antiphospholipid syndrome',
         'von Willebrand factor level': 'VWF level', 'pulmonary hypertension': 'Pulmonary hypertension'}
METHOD_STEPS = [('Samples and diagnosis', ['cohort / recruitment', 'diagnosis of CTEPH', 'controls', 'DNA source']),
                ('Genotyping and sequencing', ['genotyping platform', 'sequencing', 'HLA typing', 'imputation',
                                               'variant calling / quality control']),
                ('Analysis', ['statistical test', 'covariates', 'multiple testing', 'power', 'software']),
                ('Replication', ['replication'])]
DESIGN_GROUPS = ['Genome-wide studies (GWAS, genome sequencing, Mendelian randomisation)',
                 'Exome and gene-panel sequencing', 'Clonal haematopoiesis and somatic variants', 'HLA typing',
                 'Candidate-gene association', 'Other designs', 'Case reports']
ABBR = {'CTEPH': 'chronic thromboembolic pulmonary hypertension', 'GWAS': 'genome-wide association study',
        'L2G': 'locus-to-gene', 'PMID': 'PubMed identifier', 'CLPP': 'colocalisation posterior probability',
        'PICS': 'probabilistic identification of causal SNPs', 'EAS': 'East Asian', 'PE': 'pulmonary embolism',
        'VTE': 'venous thromboembolism', 'DVT': 'deep vein thrombosis', 'PH': 'pulmonary hypertension',
        'PAH': 'pulmonary arterial hypertension', 'β': 'effect of the variant allele as reported by the study',
        'PMC': 'PubMed Central', 'PPR': 'Europe PMC preprint record', 'HLA': 'human leukocyte antigen',
        'ClinGen': 'Clinical Genome Resource', 'PanelApp': 'Genomics England PanelApp',
        'PVOD': 'pulmonary veno-occlusive disease', 'PCH': 'pulmonary capillary haemangiomatosis',
        'HUPRA': 'hyperuricaemia, pulmonary hypertension, renal failure and alkalosis', 'VWF': 'von Willebrand factor',
        'CH': 'clonal haematopoiesis'}


# ── text helpers ─────────────────────────────────────────────────────────────
def p_parts(mant, exp):
    """P from mantissa / exponent: 7.00 × 10⁻²⁷ (no float underflow)."""
    m, e = num(pd.Series([mant]))[0], num(pd.Series([exp]))[0]
    if pd.isna(m) or pd.isna(e) or m <= 0:
        return None
    m, e = float(m), int(e)
    while m >= 9.995:
        m, e = m / 10, e + 1
    while m < 1:
        m, e = m * 10, e - 1
    if e >= -2:
        return [f'{m * 10 ** e:.2f}']
    return [f'{m:.2f} × 10', ('sup', f'−{-e}')]


def genes_parts(pairs):
    """[(symbol, score), ...] -> italic symbols with scores."""
    out = []
    for k, (g, s) in enumerate(pairs):
        if k:
            out.append(', ')
        out += [('i', g)] + ([f' {s:.2f}'] if s is not None and not pd.isna(s) else [])
    return out


def italic_list(genes, sep=', '):
    out = []
    for k, g in enumerate(genes):
        out += ([sep] if k else []) + [('i', g)]
    return out


def md_parts(text):
    """'**Lead.** text' -> rich parts with the bold runs."""
    out = []
    for k, seg in enumerate(re.split(r'\*\*(.+?)\*\*', text)):
        if seg:
            out.append(('b', seg) if k % 2 else seg)
    return out


def text_of(parts):
    return ''.join(p if isinstance(p, str) else p[1] for p in parts)


def text_len(parts):
    return len(text_of(parts))


def blank(x):
    return x is None or (isinstance(x, float) and math.isnan(x)) or (isinstance(x, str) and x == '')


def plain(v, sup=False):
    """Text of a cell read back with rich_text=True, without its superscript runs (sup=True: only those)."""
    from openpyxl.cell.rich_text import CellRichText
    if isinstance(v, CellRichText):
        return ''.join((b if isinstance(b, str) else b.text) for b in v
                       if (not isinstance(b, str) and b.font is not None and b.font.vertAlign == 'superscript') == sup)
    return '' if (v is None or sup) else str(v)


def wrap_parts(parts, width):
    """Split rich parts (str / (style, str)) into lines of <= width characters, at spaces and newlines."""
    tokens = []
    for p in parts:
        st, tx = ('', p) if isinstance(p, str) else p
        for li, line in enumerate(tx.split('\n')):
            if li:
                tokens.append(('', '\n'))
            words = line.split(' ')
            for k, wd in enumerate(words):
                tokens.append((st, wd + (' ' if k < len(words) - 1 else '')))
    lines, cur, n = [], [], 0
    for st, tx in tokens:
        if tx == '\n':
            lines.append(cur)
            cur, n = [], 0
            continue
        if tx == '':
            continue
        if n + len(tx.rstrip()) > width and cur and st not in ('sup', 'sub') and cur[-1][1].endswith(' '):
            lines.append(cur)
            cur, n = [], 0
        cur.append((st, tx))
        n += len(tx)
    if cur:
        lines.append(cur)
    out = []
    for ln in lines:
        if ln:
            ln[-1] = (ln[-1][0], ln[-1][1].rstrip())
        out.append([tx if st == '' else (st, tx) for st, tx in ln])
    return out


def year_span(a, b):
    a, b = num(pd.Series([a]))[0], num(pd.Series([b]))[0]
    if pd.isna(a):
        return None
    return f'{int(a)}' if a == b else f'{int(a)}–{int(b)}'


def fmt2(x):
    v = num(pd.Series([x]))[0]
    return None if pd.isna(v) else f'{v:.2f}'


def intx(x):
    v = num(pd.Series([x]))[0]
    return None if pd.isna(v) else f'{int(v):,}'


def count(x):
    """A count as a number cell (None when not a number)."""
    v = num(pd.Series([x]))[0]
    return None if pd.isna(v) else int(v)


def short(name):
    s = SHORT.get(name, name)
    return s if len(s) <= 30 else s[:29] + '…'


def b(x):
    return str(x) == 'True'


def cite(r):
    return ' '.join(y for y in (str(r.first_author).split()[0] if r.first_author else '', str(r.year)) if y)


def lines_needed(text, width):
    """Wrapped lines of a cell (Arial 9 pt: about 1.3 characters per unit of column width, less what wrapping at
    spaces loses)."""
    return sum(max(1, math.ceil(len(seg) / max(width * 1.3 - 4, 1))) for seg in text.split('\n'))


def plural(n, word, words=None):
    return f'{n:,} {word if n == 1 else (words or word + "s")}'


# ── the writer ───────────────────────────────────────────────────────────────
class Sheet:
    """An XlsxWriter worksheet with a format cache and a rich-text helper (as integration/gene_signals)."""

    def __init__(self, wb, ws):
        self.wb, self.ws, self._fmt = wb, ws, {}

    def fmt(self, **kw):
        props = {'font_name': FONT, 'font_size': 9, 'font_color': INK, 'valign': 'top'}
        props.update({k: v for k, v in kw.items() if v is not None})
        key = tuple(sorted(props.items()))
        if key not in self._fmt:
            self._fmt[key] = self.wb.add_format(props)
        return self._fmt[key]

    def font(self, style, b, i, color, sz):
        props = {'font_name': FONT, 'font_size': sz, 'font_color': color, 'bold': bool(b), 'italic': bool(i)}
        if style == 'i':
            props['italic'] = True
        elif style == 'b':
            props['bold'] = True
        elif style == 'bi':
            props['bold'], props['italic'] = True, True
        elif style == 'sup':
            props['font_script'] = 1
        elif style == 'g':
            props['font_color'] = INK2
        elif style == 'gi':
            props['font_color'], props['italic'] = INK2, True
        key = ('font',) + tuple(sorted(props.items()))
        if key not in self._fmt:
            self._fmt[key] = self.wb.add_format(props)
        return self._fmt[key]

    def rich(self, r, c, parts, cell, b=False, i=False, color=INK, sz=9):
        frags = []
        for p in parts:
            st, tx = ('', p) if isinstance(p, str) else p
            if tx == '':
                continue
            if frags and frags[-1][0] == st:
                frags[-1][1] += tx
            else:
                frags.append([st, tx])
        if not frags:
            self.ws.write_blank(r, c, None, self.fmt(**cell))
        elif len(frags) == 1:
            st, tx = frags[0]
            f = dict(cell, font_size=sz, font_color=INK2 if st in ('g', 'gi') else color,
                     bold=True if (b or st in ('b', 'bi')) else None,
                     italic=True if (i or st in ('i', 'bi', 'gi')) else None, font_script=1 if st == 'sup' else None)
            self.ws.write_string(r, c, tx, self.fmt(**f))
        else:
            args = []
            for st, tx in frags:
                args += [self.font(st, b, i, color, sz), tx]
            self.ws.write_rich_string(r, c, *args, self.fmt(**cell))


def new_sheet(wb, name, tab=None):
    sh = Sheet(wb, wb.add_worksheet(name))
    sh.ws.hide_gridlines(2)
    sh.ws.set_zoom(90)
    if tab:
        sh.ws.set_tab_color(tab)
    return sh


def title(sh, text, sub):
    sh.rich(0, 0, [('b', text)], {'align': 'left', 'valign': 'vcenter'}, sz=11)
    sh.rich(1, 0, sub if isinstance(sub, list) else [sub], {'align': 'left'}, i=True, color=INK2)
    sh.ws.set_row(0, 18)


def header_row(sh, top, cols):
    rot = any(c.get('rot') for c in cols)
    for j, c in enumerate(cols):
        sh.ws.set_column(j, j, c['w'])
        h = c['h'] if isinstance(c['h'], list) else [('b', c['h'])]
        cell = {'align': c.get('a', 'left'), 'valign': 'bottom', 'text_wrap': not c.get('rot'), 'bold': True,
                'top': MEDIUM, 'top_color': INK, 'bottom': THIN, 'bottom_color': INK}
        if c.get('rot'):
            cell.update({'rotation': 90, 'align': 'center'})
        sh.rich(top, j, h + [('sup', c.get('n') or '')], cell, b=True)
        if c.get('tip'):
            sh.ws.write_comment(top, j, c['tip'], {'font_name': FONT, 'font_size': 9, 'x_scale': 2})
    if rot:
        sh.ws.set_row(top, min(150, max(5.3 * max(text_len(c['h'] if isinstance(c['h'], list) else [c['h']])
                                                  for c in cols if c.get('rot')) + 12, 40)))
    else:
        hl = max(math.ceil(text_len(c['h'] if isinstance(c['h'], list) else [c['h']]) / max(c['w'] * 1.05, 1))
                 for c in cols)
        sh.ws.set_row(top, max(26, 12.5 * hl + 2))


def table(sh, top, cols, rows, links=None, freeze=1, autofilter=False):
    """A flat booktabs table. cols: [{'h': header, 'w': width, 'n': footnote letter, 'a': align, 'wrap': bool}];
    rows: lists of cells (a str, rich parts, None for '—', '' for blank); links: {(i, j): url}.
    Returns the next free row."""
    links = links or {}
    header_row(sh, top, cols)
    for i, row in enumerate(rows):
        r = top + 1 + i
        last = i == len(rows) - 1
        lines = 1
        for j, (c, v) in enumerate(zip(cols, row)):
            cell = {'align': c.get('a', 'left'), 'text_wrap': bool(c.get('wrap'))}
            if last:
                cell.update({'bottom': MEDIUM, 'bottom_color': INK})
            parts = [] if isinstance(v, str) and v == '' else [MISSING] if blank(v) else \
                (v if isinstance(v, list) else [str(v)])
            if (i, j) in links and not blank(v):
                sh.ws.write_url(r, j, links[(i, j)], sh.fmt(**cell, font_color=LINK, underline=1),
                                string=text_of(parts))
            else:
                sh.rich(r, j, parts, cell)
            if c.get('wrap'):
                lines = max(lines, lines_needed(text_of(parts), c['w']))
        if lines > 1:
            sh.ws.set_row(r, 12.0 * lines + 1)
    if not rows:
        sh.rich(top + 1, 0, ['No records.'], {'align': 'left', 'bottom': MEDIUM, 'bottom_color': INK})
    if freeze is not None:
        sh.ws.freeze_panes(top + 1, freeze)
    if autofilter and rows:
        sh.ws.autofilter(top, 0, top + len(rows), len(cols) - 1)
    return top + 1 + max(len(rows), 1)


def notes(sh, row, items, width=150):
    """items: [(letter, parts)]; letter '' for unlettered lines."""
    for letter, parts in items:
        for k, ln in enumerate(wrap_parts(parts, width - 3)):
            lead = [('sup', letter), ' '] if (letter and k == 0) else []
            sh.rich(row, 0, lead + ln, {'align': 'left', 'indent': 1 if k else None}, color=NOTE)
            row += 1
    return row


def missing_block(items):
    out = [('', [('b', 'Missing values')])]
    for where, text in items:
        out.append(('', ['   ', ('b', MISSING), f', in {where}: ', text]))
    return out


def abbrev(keys):
    return ('', ['Abbreviations: ' + '; '.join(f'{k}, {ABBR[k]}' for k in sorted(keys, key=str.lower)) + '.'])


def print_setup(ws, last_row, ncol, footer, header=3):
    ws.set_landscape()
    ws.set_paper(8)
    ws.fit_to_pages(1, 0)
    ws.repeat_rows(header, header)
    ws.print_area(0, 0, last_row, ncol - 1)
    ws.set_margins(left=0.4, right=0.4)
    ws.set_footer(f'&L{footer}&RPage &P of &N')


# ── question sheets: a brief, then a nested table ────────────────────────────
def R(lvl, kind, cells, span=None, links=None, tips=None, fills=None, hide=False, key=None, bold=None):
    """One row of a nested table. kind: group0 | group1 (a row that opens) | item | quote.
    cells: one per column (str, rich parts, number, None for '—', '' for blank); span: (j0, j1) merges those
    columns into one cell holding cells[j0]; links: {j: url or ('@', sheet, key)}; tips: {j: comment};
    fills: {j: (background, font colour)}; key: an anchor other sheets can link to; bold: a bold label (default:
    group0 rows)."""
    return {'lvl': lvl, 'kind': kind, 'cells': cells, 'span': span, 'links': links or {}, 'tips': tips or {},
            'fills': fills or {}, 'hide': hide, 'key': key, 'bold': kind == 'group0' if bold is None else bold}


def brief_lines(spec):
    """The lines above the table: [(parts, style, outline level)]."""
    width = min(int(sum(c['w'] for c in spec['cols']) * 1.05), 175)
    out = []
    if spec.get('brief'):
        out.append(([('b', 'In brief')], 'head', 0))
        for item in spec['brief']:
            for k, ln in enumerate(wrap_parts(item, width - 3)):
                out.append(((['• '] if k == 0 else ['   ']) + ln, 'text', 0))
    for head, paras in spec.get('summaries') or []:
        out.append(([('b', head)], 'shead', 0))
        for p in paras:
            for ln in wrap_parts(p, width):
                out.append((ln, 'text', 1))
            out.append(([], 'text', 1))
    return out + [([], 'text', 0)] if out else out


def layout(spec):
    """Row of the header and of every body row (0-based)."""
    top = 3 + len(brief_lines(spec)) + (1 if spec.get('spanners') else 0)
    return top, {r['key']: top + 1 + i for i, r in enumerate(spec['rows']) if r['key'] is not None}


STYLE = {'group0': {'bg_color': GROUP0}, 'group1': {'bg_color': GROUP1}, 'item': {}, 'quote': {}}


def write_spec(wb, spec, anchors, foot):
    sh = new_sheet(wb, spec['name'], spec.get('tab'))
    ws, cols = sh.ws, spec['cols']
    ws.outline_settings(True, False, True, False)
    title(sh, spec['title'], spec['sub'])
    if spec['name'] != 'Guide':
        ws.write_url(0, len(cols) - 1, "internal:'Guide'!A1", sh.fmt(font_color=LINK, underline=1, align='right'),
                     string='← Guide')
    r = 3
    for parts, style, lvl in brief_lines(spec):
        sh.rich(r, 0, parts, {'align': 'left'}, color=INK if style != 'shead' else LINK)
        if lvl or style == 'shead':
            ws.set_row(r, None, None, {'level': lvl})
        r += 1
    top, _ = layout(spec)
    if spec.get('spanners'):
        for j0, j1, text in spec['spanners']:
            f = sh.fmt(bold=True, align='center', valign='bottom', text_wrap=True, bottom=THIN, bottom_color=INK)
            if j1 > j0:
                ws.merge_range(top - 1, j0, top - 1, j1, text, f)
            else:
                ws.write_string(top - 1, j0, text, f)
        ws.set_row(top - 1, 24)
    header_row(sh, top, cols)
    rows = spec['rows']
    for i, row in enumerate(rows):
        rr = top + 1 + i
        last = i == len(rows) - 1
        style = STYLE[row['kind']]
        height = 1
        j = 0
        while j < len(cols):
            j1 = row['span'][1] if row['span'] and row['span'][0] == j else j
            c, v = cols[j], row['cells'][j]
            cell = {'align': c.get('a', 'left') if j1 == j else 'left', 'text_wrap': bool(c.get('wrap') or j1 > j),
                    **style}
            if j == 0:
                cell['indent'] = row['lvl']
            if last:
                cell.update({'bottom': MEDIUM, 'bottom_color': INK})
            color = INK2 if row['kind'] == 'quote' else INK
            bold = j == 0 and row['bold']
            if j in row['fills']:
                cell['bg_color'], color = row['fills'][j]
                bold = color != INK
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                ws.write_number(rr, j, v, sh.fmt(**cell, num_format='#,##0', font_color=color))
                text = f'{v:,}'
            else:
                parts = [] if isinstance(v, str) and v == '' else [MISSING] if blank(v) else \
                    (v if isinstance(v, list) else [str(v)])
                text = text_of(parts)
                if j1 > j:
                    ws.merge_range(rr, j, rr, j1, '', sh.fmt(**cell))
                link = row['links'].get(j)
                if link and not blank(v):
                    url = link if isinstance(link, str) else \
                        f"internal:'{link[1]}'!A{anchors[(link[1], link[2])] + 1}"
                    ital = all(not isinstance(x, str) and x[0] in ('i', 'bi') for x in parts)   # a gene symbol
                    ws.write_url(rr, j, url, sh.fmt(**cell, font_color=LINK, underline=1, bold=bold or None,
                                                    italic=ital or None), string=text)
                else:
                    sh.rich(rr, j, parts, cell, b=bold, color=color)
            if j in row['tips']:
                tip = row['tips'][j]
                ws.write_comment(rr, j, tip, {'font_name': FONT, 'font_size': 9, 'x_scale': 3.2,
                                              'y_scale': max(1.0, lines_needed(tip, 60) / 4.5)})
            if cell['text_wrap']:
                height = max(height, lines_needed(text, sum(cols[k]['w'] for k in range(j, j1 + 1))))
            j = j1 + 1
        nxt = rows[i + 1] if i + 1 < len(rows) else None
        opts = {'level': row['lvl'], 'hidden': row['hide'],
                'collapsed': bool(nxt and nxt['lvl'] > row['lvl'] and nxt['hide'])}
        ws.set_row(rr, 12.0 * height + 1 if height > 1 else None, None, opts)
    end = top + 1 + max(len(rows), 1)
    if not rows:
        sh.rich(top + 1, 0, ['No records.'], {'align': 'left', 'bottom': MEDIUM, 'bottom_color': INK})
    if spec.get('freeze'):
        ws.freeze_panes(top + 1, spec['freeze'])
    end = notes(sh, end + 1, spec.get('notes') or [], spec.get('note_width', 150))
    print_setup(ws, end, len(cols), foot, header=top)
    return {'header': top, 'rows': len(rows), 'spanner': bool(spec.get('spanners'))}


def spec(name, title_, sub, cols, rows, **kw):
    return {'name': name, 'title': title_, 'sub': sub, 'cols': cols, 'rows': rows, 'tab': '#1F4E79', **kw}


def summary_block(D, kind, head):
    s = D['summ'][D['summ'].kind == kind]
    if not len(s):
        return []
    x = s.iloc[0]
    if not b(x.grounded):
        return [(f'{head}: withheld, it failed its check', [])]
    return [(f'{head} (generated by a language model from the records below; it cites only their PMIDs)',
             [md_parts(p) for p in x.summary.split('\n') if p.strip()])]


# ── data shared by several sheets ────────────────────────────────────────────
def prepare(D):
    cond, gc = D['conditions'], D['gene_condition'].copy()
    gc['code'] = gc.evidence.map(KIND)
    D['gc'] = gc
    D['cteph_genes'] = set(gc[gc.node_id == CORE].symbol)
    D['cteph_kinds'] = gc[gc.node_id == CORE].groupby('symbol').code.apply(lambda x: ''.join(sorted(set(x),
                                                                           key='GCL'.index))).to_dict()
    D['rel_n'] = gc[gc.node_id != CORE].groupby('symbol').node_id.nunique().to_dict()
    cst = D['cst'].copy()
    cst['population'] = [study_population(r.country, r.ancestry) for r in cst.itertuples()]
    cst['jp'] = [is_japanese(r) for r in cst.itertuples()]
    cst['eas'] = [is_east_asian(r) for r in cst.itertuples()]
    cst['cite'] = [f'{r.first_author} {r.year}' for r in cst.itertuples()]
    D['cst'] = cst
    fx = assign_themes(D['cfi'], D['themes']).merge(
        cst[['pmid', 'year', 'cite', 'design', 'population', 'jp', 'eas', 'n_cases', 'n_controls']], on='pmid',
        how='left')
    D['fx'] = fx.assign(_y=num(fx.year)).sort_values(['_y', 'pmid'])
    D['gt'] = gene_tallies(D['fx'])
    other = sorted(set(D['gt'][D['gt'].theme == OTHER_GENES].label))
    if other:
        log(f'genes in no theme (listed under "{OTHER_GENES}"): {", ".join(other)}')
    D['in_sheet5'] = set(cond[cond.genetics.eq('full') & cond.ot_ids.ne('')].node_id)


def rel_category(D, node):
    c = D['conditions'].set_index('node_id')
    cat = c.category[node]
    if cat != 'Shared genetic signal':
        return cat
    return 'multi' if int(c.coloc_loci[node]) >= 2 else 'single'


# ── Guide ────────────────────────────────────────────────────────────────────
def answers(D):
    st, cond, gc = D['tidy'], D['conditions'], D['gc']
    bs = st['conditions_by_source']
    m2 = cond[cond.coloc_loci.map(lambda x: int(x or 0)) >= 2]
    multi = m2[m2.category != 'Shared genetic signal'].name.tolist()
    n_other = int((m2.category == 'Shared genetic signal').sum())
    shared = sorted(((D['rel_n'].get(g, 0), g) for g in D['cteph_genes'] if D['rel_n'].get(g, 0)), reverse=True)
    gt = D['gt']
    rep = gt[(gt.associated >= 2) & (gt.theme != NO_GENE)].sort_values('associated', ascending=False)
    cst = D['cst']
    des = cst.design.str.split('; ').explode().value_counts()
    mech = D['summ'][D['summ'].kind == 'mechanism'].assign(n=lambda d: num(d.n_targets)).sort_values('n',
                                                                                                   ascending=False)
    eas_rel = int(D['relations'].east_asian.map(b).sum())
    full = cond[cond.node_id.isin(D['in_sheet5'])]
    ch = D['changes_json']
    return {
        '1 Related conditions': [f'{st["conditions"]} conditions: {bs["clinical"]} with clinical evidence quoted from '
                                 f'PubMed, {bs["ontology"]} by the disease ontology, {bs["colocalisation"]} sharing a '
                                 'GWAS signal with a CTEPH locus. Two or more CTEPH loci are shared by '
                                 + ', '.join(short(n) for n in multi) + f', and by {n_other} traits without a '
                                 'clinical relation.'],
        '2 Shared genes': [f'{len(D["cteph_genes"])} genes have genetic evidence in CTEPH; {len(shared)} of them also '
                           'in a related condition, most widely '] +
        italic_list([f'{g}' for _n, g in shared[:6]]) + ['.'],
        '3 CTEPH genetics': [f'{len(cst)} genetic studies of CTEPH ({st["cteph_findings"]} findings). Associated in two '
                             'or more studies: '] +
        italic_list([f'{r.label}' for r in rep.itertuples()]) +
        ['. One GWAS in Open Targets, with loci at ', ('i', 'ABO'), ', ', ('i', 'F11'), ' and ', ('i', 'FGA'), '/',
         ('i', 'FGG'), '.'],
        '4 Study methods': ['Designs: ' + ', '.join(f'{k} {v}' for k, v in des.head(4).items()) + '; methods from the full text for '
                            f'{int(cst.methods_source.eq("full text").sum())} studies, from the abstract for '
                            f'{int(cst.methods_source.eq("abstract").sum())}.'],
        '5 Condition genetics': [f'{len(full)} conditions with genetic data: {st["gwas_studies"]} GWAS studies '
                                 f'({st["eas_studies"]} East Asian), {st["credible_sets"]:,} credible sets, '
                                 f'{st["mendelian_links"]} curated gene links.'],
        '6 Mechanisms': [f'{len(mech)} Reactome themes of the text-mined CTEPH targets; the largest: ' +
                         ', '.join(f'{r.title} ({int(r.n)})' for r in mech.head(4).itertuples()) + '.'],
        '7 East Asian': [f'{eas_rel} clinical evidence items, {int(cst.eas.sum())} genetic studies of CTEPH '
                         f'({int(cst.jp.sum())} Japanese) and {st["eas_studies"]} GWAS of related conditions.'],
        '8 Changes': ['First snapshot: nothing to compare with.' if ch['previous'] is None else
                      f'Since {ch["previous"]}: ' + ('; '.join(f'{k.replace("|", " ")} {v}' for k, v in
                                                           ch['counts'].items()) or 'no change') + '.'],
    }


def s_guide(wb, D, foot):
    sh = new_sheet(wb, 'Guide', '#1A1A1A')
    st, ch = D['tidy'], D['changes_json']
    title(sh, 'CTEPH knowledge graph: conditions related to CTEPH and their genetics',
          f'Open Targets Platform release {st["release"]} and PubMed, fetched {st["fetch_date"]}. External knowledge '
          'only. Each sheet answers one question; open its rows for the evidence.')
    ans = answers(D)
    rows, links = [], {}
    for i, (s, q) in enumerate(QUESTIONS):
        rows.append([s, q, ans[s]])
        links[(i, 0)] = f"internal:'{s}'!A1"
    r = table(sh, 3, [{'h': 'Sheet', 'w': 22}, {'h': 'Question', 'w': 44, 'wrap': True},
                      {'h': 'Answer in brief', 'w': 100, 'wrap': True}], rows, links, freeze=None)
    rows, links = [], {}
    for i, (s, d) in enumerate(APPENDIX):
        rows.append([s, d, f'{D["appendix_rows"][s]:,}'])
        links[(i, 0)] = f"internal:'{s}'!A1"
    r = table(sh, r + 1, [{'h': 'Appendix', 'w': 22}, {'h': 'Flat table, with filters', 'w': 44, 'wrap': True},
                          {'h': 'Records', 'w': 100}], rows, links, freeze=None)
    items = [('', [('b', 'How to read')]),
             ('', ['   Rows are nested: group > record > source. Open or close them with + / − at the left, or show a '
                   'whole level with the buttons 1 2 3 above the row numbers.']),
             ('', ['   Shaded bold rows are groups; shaded rows open to their records; grey rows hold source quotes '
                   'and sentences. Blue underlined text is a link: to PubMed, or to the same gene or condition on '
                   'another sheet.']),
             ('', ['   A cell with a red corner has a note: hover over it for the source passage or the evidence '
                   'behind a code.']),
             ('', ['   Gene map codes: G, GWAS (the top locus-to-gene gene of a credible set); C, curated gene '
                   'evidence (ClinGen, PanelApp, Orphanet, ClinVar, UniProt, Gene2Phenotype, gene burden); L, a '
                   'genetic study of CTEPH in the literature. Fill: blue G, green C, red L, grey two kinds.']),
             ('', ['   Paragraphs headed "Generated summary" were written by a language model from the records '
                   'listed with them, and every PMID and gene they cite was checked; they are not a systematic '
                   'review.']),
             ('', ['   ', ('b', MISSING), ' marks a value that is not available; the reason is given under each '
                   'table.'])]
    r = notes(sh, r + 1, items, 168)
    srows = [['Open Targets release', st['release'], ''], ['Fetched', st['fetch_date'], ''],
             ['Previous snapshot', ch['previous'] or 'none', ''] if ch['previous'] else
             ['Previous snapshot', 'none', 'this is the first'],
             ['Related conditions', f'{st["conditions"]}', ''], ['GWAS studies / credible sets',
                                                                 f'{st["gwas_studies"]} / {st["credible_sets"]:,}', ''],
             [f'{CURATED} gene links', f'{st["mendelian_links"]}', ''],
             ['PubMed records screened / genetic studies of CTEPH',
              f'{st["cteph_genetic_records"]} / {st["cteph_genetic_studies"]}', ''],
             ['Knowledge-graph nodes / edges', f'{st["graph_nodes"]:,} / {st["graph_edges"]:,}', '']]
    r = table(sh, r + 1, [{'h': 'Snapshot', 'w': 22}, {'h': 'Value', 'w': 44}, {'h': 'Note', 'w': 100}], srows,
              freeze=None)
    items = [('', [('b', 'Caveats')])] + [('', ['   ' + c]) for c in D['caveats']]
    items.append(abbrev(['CTEPH', 'GWAS', 'PMID']))
    r = notes(sh, r + 1, items, 168)
    print_setup(sh.ws, r, 3, foot)
    return {'header': 3, 'rows': len(QUESTIONS), 'spanner': False}


# ── 1 related conditions ─────────────────────────────────────────────────────
def spec_conditions(D):
    cond, rel = D['conditions'], D['relations'].copy()
    gc = D['gc']
    gene_sets = gc.groupby('node_id').symbol.apply(set).to_dict()
    co = D['coloc']
    loc_gene = dict(zip(co.cteph_locus, co.cteph_gene))
    cols = [{'h': 'Category › condition › evidence', 'w': 44, 'wrap': True},
            {'h': 'Clinical items', 'w': 8, 'a': 'center', 'n': 'a'},
            {'h': 'Ontology', 'w': 8, 'a': 'center', 'n': 'b'},
            {'h': 'CTEPH loci shared', 'w': 13, 'n': 'c'},
            {'h': 'East Asian evidence', 'w': 9, 'a': 'center'},
            {'h': 'GWAS studies', 'w': 8, 'a': 'right'},
            {'h': 'EAS studies', 'w': 7, 'a': 'right', 'n': 'd'},
            {'h': f'{CURATED} genes', 'w': 8, 'a': 'center', 'n': 'e'},
            {'h': 'Genes shared with CTEPH', 'w': 26, 'wrap': True, 'n': 'f'},
            {'h': 'Link', 'w': 14, 'n': 'g'}]
    nc = len(cols)
    rel['_t'] = rel.evidence_type.map({'clinical': 0, 'ontology': 1, 'colocalisation': 2})
    groups = CATEGORY_ORDER + ['multi', 'single']
    cond = cond.assign(_g=[rel_category(D, n) for n in cond.node_id])
    rows = []
    for g in groups:
        sub = cond[cond._g == g].sort_values('name', key=lambda s: s.str.lower())
        if not len(sub):
            continue
        label = g if g in CATEGORY_ORDER else ('Shared GWAS signal: two or more CTEPH loci' if g == 'multi' else
                                               'Shared GWAS signal: one CTEPH locus')
        rows.append(R(0, 'group0', [f'{label} ({len(sub)})', CATEGORY_DESC[g]] + [''] * (nc - 2), span=(1, nc - 1)))
        for c in sub.itertuples():
            ev = rel[rel.node_id == c.node_id].sort_values(['_t', 'pmid'])
            loci = sorted({loc_gene.get(x, x) for v in ev.loci for x in v.split() if x})
            ea = ev.east_asian.map(b).any()
            shared = sorted(gene_sets.get(c.node_id, set()) & D['cteph_genes'])
            full = c.genetics == 'full'
            link = ('@', '5 Condition genetics', c.node_id) if c.node_id in D['in_sheet5'] else None
            cells = [c.name, count(c.n_clinical) or '', 'yes' if 'ontology' in c.sources else '',
                     italic_list(loci) if loci else '', 'yes' if ea else '',
                     count(c.n_gwas_studies) if full else '', count(c.n_eas_studies) if full else '',
                     count(c.n_mendelian_genes) if full else '', italic_list(shared) if shared else '',
                     'genetics →' if link else '']
            rows.append(R(1, 'group1', cells, links={nc - 1: link} if link else None,
                          hide=g == 'single', key=('1 Related conditions', c.node_id)))
            for x in ev.itertuples():
                if x.evidence_type == 'clinical':
                    lab = f'Clinical: {x.relation}' if x.relation else 'Clinical'
                    txt = [x.detail] + (['\n', ('gi', f'“{x.quote}”')] if x.quote else []) + \
                        ([('g', '  (' + '; '.join(y for y in (x.design, x.population) if y) + ')')]
                         if (x.design or x.population) else [])
                elif x.evidence_type == 'ontology':
                    lab, txt = 'Ontology', [x.detail]
                else:
                    lab = 'Shared GWAS signal'
                    n = int(x.n_studies or 0)
                    txt = [f'{n} GWAS {"signal colocalises" if n == 1 else "signals colocalise"} with the CTEPH '
                           f'{"locus" if len(x.loci.split()) == 1 else "loci"} ']
                    for k, v in enumerate(x.loci.split()):
                        txt += ([', '] if k else []) + [('i', loc_gene.get(v, v)), f' ({v})']
                    txt += [f'; strongest CLPP {float(x.strength):.2f}.' if x.strength else '.']
                ref = (f'{x.citation.split(",")[0]}' if x.pmid else x.citation) or None
                rows.append(R(2, 'quote', [lab] + [txt] + [''] * (nc - 3) + [ref], span=(1, nc - 2),
                              links={nc - 1: pub_url(x.pmid)} if x.pmid else None, hide=True))
    bs = D['tidy']['conditions_by_source']
    brief = [[f'{len(cond)} conditions: {bs["clinical"]} with clinical evidence (quotes from PubMed abstracts, '
              f'checked at every build), {bs["ontology"]} by the disease ontology, {bs["colocalisation"]} sharing a '
              'GWAS signal with a CTEPH locus. A condition can have several kinds.'],
             ['The precursor / spectrum, risk-factor, comorbidity, intermediate-trait and PH groups are open; the '
              'traits sharing one CTEPH locus only are closed (level 1 button to see every condition).'],
             ['Contrary findings are kept: against survivors of pulmonary embolism, most classic risk factors were '
              'not enriched (open Risk factor and Comorbidity).']]
    nts = [('a', ['Clinical evidence items: guideline, review or study sentences quoted from PubMed abstracts.']),
           ('b', ['A relation in the disease ontology: the parent of CTEPH, its siblings, or a child of pulmonary '
                  'arterial hypertension.']),
           ('c', ['CTEPH GWAS loci, by their locus-to-gene gene, whose credible set colocalises with a GWAS signal '
                  'of the condition (eCAVIAR CLPP ≥ 0.01).']),
           ('d', [f'GWAS studies with a mostly East Asian sample (≥ {D["eas_min"]:.0%}).']),
           ('e', ['Genes with curated gene-disease evidence (ClinGen, PanelApp, Orphanet, ClinVar, UniProt, '
                  'Gene2Phenotype, gene burden) for the condition or its descendants.']),
           ('f', ['Genes with genetic evidence in both the condition and CTEPH (sheet 2).']),
           ('g', ['Condition rows: the condition on sheet 5; evidence rows: the PubMed record.'])]
    nts += [('', [('b', 'Missing values')]),
            ('', ['   Empty cells: no evidence of that kind, or no genetic data fetched (a procedure, a phenotype too '
                  'broad, or a trait taken only for its shared signal).'])]
    nts += missing_block([('Link', 'no PubMed record (ontology and colocalisation evidence)')])
    nts.append(abbrev(['CTEPH', 'GWAS', 'EAS', 'CLPP', 'PH']))
    return spec('1 Related conditions', '1  ' + QUESTIONS[0][1],
                'Category › condition › evidence. Open a condition (+) for its quotes, ontology relation and shared '
                'GWAS signals.', cols, rows, brief=brief, notes=nts)


# ── 2 shared genes ───────────────────────────────────────────────────────────
def spec_genes(D):
    gc, cond = D['gc'], D['conditions']
    nodes = [CORE] + [n for cat in CATEGORY_ORDER for n in
                      cond[(cond.category == cat) & cond.node_id.isin(set(gc.node_id))]
                      .sort_values('name', key=lambda s: s.str.lower()).node_id]
    names = dict(zip(cond.node_id, cond.name))
    names[CORE] = 'chronic thromboembolic pulmonary hypertension'
    cell = {}
    for (g, n), x in gc.groupby(['symbol', 'node_id']):
        codes = ''.join(sorted(set(x.code), key='GCL'.index))
        tip = '\n'.join(f'{dict(G="GWAS", C=CURATED, L="CTEPH literature")[k]}: {d}'
                        for k, d in zip(x.code, x.details))
        cell[(g, n)] = (codes, f'{x.symbol.iloc[0]} in {names.get(n, n)}\n{tip}')
    rel_n, cteph = D['rel_n'], D['cteph_genes']
    lit_genes = set(D['gt'].symbol)
    groups = [('Genes with genetic evidence in CTEPH', sorted(cteph, key=lambda g: (-rel_n.get(g, 0), g)), False),
              ('Genes in three or more related conditions', sorted([g for g, k in rel_n.items() if k >= 3 and g not in
                                                                    cteph], key=lambda g: (-rel_n[g], g)), False),
              ('Genes in two related conditions', sorted(g for g, k in rel_n.items() if k == 2 and g not in cteph),
               True)]
    cols = [{'h': 'Gene', 'w': 12}, {'h': 'Related conditions', 'w': 6, 'a': 'right', 'rot': True, 'n': 'a'}] + \
        [{'h': short(names.get(n, n)), 'w': 4.3, 'a': 'center', 'rot': True,
          'tip': names.get(n, n)} for n in nodes]
    rows = []
    for label, genes, closed in groups:
        rows.append(R(0, 'group0', [f'{label} ({len(genes)})'] + [''] * (len(cols) - 1), span=(0, len(cols) - 1)))
        for g in genes:
            cells, fills, tips = [[('i', g)], rel_n.get(g, 0)], {}, {}
            for k, n in enumerate(nodes, start=2):
                codes, tip = cell.get((g, n), ('', ''))
                cells.append(codes)
                if codes:
                    fills[k] = (MULTI_FILL, '#FFFFFF') if len(codes) > 1 else (KIND_FILL[codes], INK)
                    tips[k] = tip
            links = {0: ('@', '3 CTEPH genetics', g)} if g in lit_genes else None
            rows.append(R(1, 'item', cells, links=links, tips=tips, fills=fills, hide=closed,
                          key=('2 Shared genes', g)))
    spans, j = [], 3
    for cat in CATEGORY_ORDER:
        k = sum(1 for n in nodes[1:] if cond.set_index('node_id').category[n] == cat)
        if k:
            spans.append((j, j + k - 1, cat))
            j += k
    shared = sorted(((rel_n.get(g, 0), g) for g in cteph if rel_n.get(g, 0)), reverse=True)
    brief = [[f'{len(cteph)} genes have genetic evidence in CTEPH; {len(shared)} of them also in a related condition, '
              'most widely '] + italic_list([f'{g} ({n})' for n, g in shared[:8]]) + ['.'],
             ['Codes: G GWAS (blue), C curated (green), L CTEPH literature (red); grey, two kinds. Hover over a code '
              'for its evidence; a linked gene opens its CTEPH findings (sheet 3).']]
    nts = [('a', ['Related conditions (not counting CTEPH) in which the gene has genetic evidence.']),
           ('', ['G: the gene is the top locus-to-gene gene (score ≥ ' f'{D["l2g_min"]}) of a credible set of the '
                 'condition. C: curated gene evidence for the condition or its descendants. L: a genetic study of '
                 'CTEPH reports an association, an enrichment or a variant in the gene.']),
           ('', ['Columns are the conditions with genetic data, by category; hover over a column header for its full '
                 'name.'])]
    nts.append(abbrev(['CTEPH', 'GWAS', 'L2G', 'PAH', 'PVOD', 'PCH', 'HUPRA', 'VWF', 'VTE', 'DVT', 'PE']))
    return spec('2 Shared genes', '2  ' + QUESTIONS[1][1],
                'Gene × condition map of genetic evidence: CTEPH first, then the related conditions by category.',
                cols, rows, brief=brief, notes=nts, spanners=spans, freeze=3)


# ── 3 CTEPH genetics ─────────────────────────────────────────────────────────
def spec_cteph(D):
    fx, gt, themes = D['fx'], D['gt'], D['themes']
    cols = [{'h': 'Theme › gene › study', 'w': 30, 'wrap': True}, {'h': 'Design', 'w': 18, 'wrap': True},
            {'h': 'Population', 'w': 11, 'wrap': True}, {'h': 'Cases', 'w': 7, 'a': 'right'},
            {'h': 'Controls', 'w': 8, 'a': 'right'}, {'h': 'Variant', 'w': 16, 'wrap': True},
            {'h': 'Result', 'w': 11, 'n': 'a'}, {'h': 'Effect', 'w': 18, 'wrap': True},
            {'h': 'P', 'w': 10, 'wrap': True}, {'h': 'Source sentence', 'w': 60, 'wrap': True, 'n': 'b'},
            {'h': 'Link', 'w': 13}]
    nc = len(cols)
    rows = []
    # the CTEPH GWAS in Open Targets
    lo, st, co = D['loci'], D['studies'], D['coloc']
    cl, cs = lo[lo.node_id == CORE], st[st.node_id == CORE]
    if len(cs):
        s0 = cs.iloc[0]
        head = f'{cite(s0)}: {intx(s0.n_cases)} cases and {intx(s0.n_controls)} controls, ' \
               f'{ancestry_label(s0.ancestry)}. Each locus colocalises with GWAS signals of venous thrombosis and ' \
               'blood traits (sheet 1, A3).'
        rows.append(R(0, 'group0', [f'Genome-wide loci (Open Targets GWAS, {len(cl)})', head] + [''] * (nc - 2),
                      span=(1, nc - 1)))
        for x in cl.assign(_p=num(cl.log10_p)).sort_values('_p').itertuples():
            cc = co[co.cteph_locus == x.rsid]
            traits = cc.condition.replace('', pd.NA).dropna().value_counts()
            l2g = [(getattr(x, f'l2g{k}_symbol'), num(pd.Series([getattr(x, f'l2g{k}_score')]))[0]) for k in (1, 2, 3)
                   if getattr(x, f'l2g{k}_symbol')]
            txt = ['P ', *(p_parts(x.p_mantissa, x.p_exponent) or [MISSING]), f' ({variant_label(x.variant_id)}); '
                   'locus-to-gene ', *genes_parts(l2g), f'; colocalises with {len(cc)} GWAS signals, most often ' +
                   ', '.join(f'{t} ({n})' for t, n in traits.head(5).items()) + '.']
            rows.append(R(1, 'group1', [[x.rsid + ' ', ('i', x.l2g1_symbol)], txt] + [''] * (nc - 3) +
                          [f'PMID {s0.pmid}'], span=(1, nc - 2), links={nc - 1: pub_url(s0.pmid)}))
    # the literature, by theme
    desc = {t['name']: t['description'] for t in themes}
    desc.update({OTHER_GENES: 'Genes no theme lists yet.', NO_GENE: 'Findings that name no gene (variants, gene '
                 'sets or study-level results).'})
    for th in theme_order(themes):
        g = gt[gt.theme == th].sort_values(['associated', 'n_studies', 'label'], ascending=[False, False, True])
        if not len(g):
            continue
        f = fx[fx.theme == th]
        rep = g[(g.associated >= 2) & (g.theme != NO_GENE)]
        ea = f[f.eas.eq(True)].pmid.nunique()
        txt = [desc[th] + f' {len(g)} genes, {f.pmid.nunique()} studies, {len(f)} findings.']
        if len(rep):
            txt += [' Associated in two or more studies: '] + italic_list([f'{r.label} ({r.associated})'
                                                                         for r in rep.itertuples()]) + ['.']
        if ea:
            txt += [f' East Asian studies: {ea}.']
        rows.append(R(0, 'group0', [th, txt] + [''] * (nc - 2), span=(1, nc - 1)))
        for x in g.itertuples():
            ff = f[f.label == x.label]
            n_rel = D['rel_n'].get(x.symbol, 0) if x.symbol else 0
            n_ea = int(ff[ff.eas.eq(True)].pmid.nunique())
            note = tally_text(x) + f' ({x.n_studies} {"study" if x.n_studies == 1 else "studies"}' + \
                (f', {n_ea} East Asian' if n_ea else '') + ').'
            if n_rel:
                note += f' Genetic evidence in {n_rel} related {"condition" if n_rel == 1 else "conditions"} too.'
            link = ('@', '2 Shared genes', x.symbol) if x.symbol in D['gene_map_genes'] else None
            rows.append(R(1, 'group1', [[('bi', x.label)], note] + [''] * (nc - 3) + ['map →' if link else ''],
                          span=(1, nc - 2), links={nc - 1: link} if link else None,
                          key=('3 CTEPH genetics', x.label) if x.symbol else None))
            for y in ff.itertuples():
                rows.append(R(2, 'quote', [y.cite or f'PMID {y.pmid}', y.design or None, y.population or None,
                                           count(y.n_cases) if y.n_cases else '', count(y.n_controls)
                                           if y.n_controls else '', y.variant or '', y.result, y.effect or '',
                                           y.p_value or '', [('gi', y.quote)], f'PMID {y.pmid}'],
                              links={nc - 1: pub_url(y.pmid)}, hide=True))
    cst = D['cst']
    rep = gt[(gt.associated >= 2) & (gt.theme != NO_GENE)].sort_values('associated', ascending=False)
    brief = [[f'{len(cst)} genetic studies of CTEPH in PubMed ({D["tidy"]["cteph_genetic_records"]} records screened), '
              f'{len(fx)} findings on {gt.label.nunique()} genes or gene sets. Associated in two or more studies: '] +
             italic_list([f'{r.label} ({r.associated})' for r in rep.itertuples()]) + ['.'],
             ['The only CTEPH GWAS in Open Targets has three loci: ', ('i', 'ABO'), ', ', ('i', 'F11'), ' and ',
              ('i', 'FGA'), '/', ('i', 'FGG'), '. The literature is grouped by gene theme: open a gene (+) for each '
              'finding with its source sentence.']]
    nts = [('a', ['As the study reports it: associated, not associated, enriched (a group of variants), identified (a '
                  'variant found, without a test), or other.']),
           ('b', ['Verbatim from the abstract or the open-access full text; every gene, variant, number and quote was '
                  'checked against it.']),
           ('', ['Gene themes are an editorial grouping, reviewed by hand. "Associated in two or more studies" counts '
                 'studies, not independent samples: some share cases.'])]
    nts += [('', [('b', 'Missing values')]),
            ('', ['   Empty cells: not stated in the source sentence or the abstract.'])]
    nts += missing_block([('Design, Population', 'not stated in the abstract'),
                          ('P', 'the locus has no P (not expected)')])
    nts.append(abbrev(['CTEPH', 'GWAS', 'PMID', 'HLA', 'CH']))
    return spec('3 CTEPH genetics', '3  ' + QUESTIONS[2][1],
                'The CTEPH GWAS loci, then the genetic studies in PubMed by gene theme: theme › gene › finding.',
                cols, rows, brief=brief, summaries=summary_block(D, 'findings', 'Generated summary of the findings'),
                notes=nts)


# ── 4 study methods ──────────────────────────────────────────────────────────
def design_group(s):
    d = s.design
    if s.study_type == 'case report':
        return DESIGN_GROUPS[6]
    if any(k in d for k in ('GWAS', 'whole-genome', 'Mendelian randomization')):
        return DESIGN_GROUPS[0]
    if 'exome' in d or 'panel' in d:
        return DESIGN_GROUPS[1]
    if 'clonal' in d:
        return DESIGN_GROUPS[2]
    if 'HLA' in d:
        return DESIGN_GROUPS[3]
    if 'candidate' in d:
        return DESIGN_GROUPS[4]
    return DESIGN_GROUPS[5]


def spec_methods(D):
    cst, cmf = D['cst'].copy(), D['cmf']
    known = {f for _s, fs in METHOD_STEPS for f in fs}
    cols = [{'h': 'Design › study', 'w': 24, 'wrap': True}, {'h': 'Population', 'w': 11, 'wrap': True},
            {'h': 'Cases', 'w': 7, 'a': 'right'}, {'h': 'Controls', 'w': 8, 'a': 'right'},
            {'h': 'Methods from', 'w': 9, 'a': 'center', 'n': 'a'}] + \
        [{'h': s, 'w': 20 if s == 'Replication' else 40, 'wrap': True, 'n': 'b' if k == 0 else None}
         for k, (s, _f) in enumerate(METHOD_STEPS)] + [{'h': 'Link', 'w': 13}]
    nc = len(cols)
    cst['group'] = [design_group(r) for r in cst.itertuples()]
    cst['_y'] = num(cst.year)
    rows = []
    for gname in DESIGN_GROUPS:
        sub = cst[cst.group == gname].sort_values(['_y', 'pmid'])
        if not len(sub):
            continue
        rows.append(R(0, 'group0', [f'{gname} ({len(sub)})'] + [''] * (nc - 1), span=(0, nc - 1)))
        for s in sub.itertuples():
            m = cmf[cmf.pmid == s.pmid]
            cells, tips = [], {}
            for k, (step, fields) in enumerate(METHOD_STEPS):
                fs = fields + ([f for f in m.field.unique() if f not in known] if step == 'Analysis' else [])
                mm = m[m.field.isin(fs)].assign(_o=lambda d: d.field.map(lambda f: fs.index(f)))
                mm = mm.sort_values('_o')
                parts, tip = [], []
                for y in mm.itertuples():
                    parts += (['\n'] if parts else []) + ([('b', f'{y.field[0].upper() + y.field[1:]}: ')]
                                                          if step != 'Replication' else []) + [y.value]
                    tip.append(f'{y.field}: “{y.quote}”')
                cells.append(parts or None)
                if tip:
                    tips[5 + k] = f'Source ({s.methods_source}):\n' + '\n\n'.join(tip)
            src = s.methods_source or None
            rows.append(R(1, 'item', [s.cite, s.population,
                                      count(s.n_cases) if s.n_cases else '', count(s.n_controls) if s.n_controls
                                      else '', src] + cells + [f'PMID {s.pmid}'],
                          links={nc - 1: pub_url(s.pmid)}, tips=tips,
                          fills={4: (SOURCE_FILL[src], INK)} if src in SOURCE_FILL else None))
    des = cst.design.str.split('; ').explode().value_counts()
    pops = cst.population.value_counts()
    nca = num(cst.n_cases)
    gw = D['studies']
    brief = [['Designs (a study can have several): ' + ', '.join(f'{k} {v}' for k, v in des.items()) + '.'],
             ['Populations: ' + ', '.join(f'{k} {v}' for k, v in pops.items()) + '. Cases per study: median '
              f'{nca.median():.0f}, range {nca.min():.0f}–{nca.max():,.0f} (where stated).'],
             [f'Methods were read from the open-access Methods section for {int(cst.methods_source.eq("full text").sum())}'
              f' studies and from the abstract for {int(cst.methods_source.eq("abstract").sum())} (no open-access full '
              'text). Hover over a methods cell for its source passages.'],
             [f'GWAS in Open Targets (CTEPH and the related conditions): {len(gw)} studies; fine-mapping of their '
              'credible sets: ' + ', '.join(f'{k} {v:,}' for k, v in D['loci'].finemapping.replace('', 'not given')
                                            .value_counts().items()) + '.']]
    nts = [('a', ['Full text: the Methods section of the open-access article (Europe PMC); abstract: the PubMed '
                  'abstract, used when no full text could be fetched.']),
           ('b', ['Methods as the source states them, by step; every value was checked against its source passage '
                  '(in the cell note).'])]
    nts += [('', [('b', 'Missing values')]),
            ('', ['   Empty cells: the number of cases or controls is not stated.'])]
    nts += missing_block([('the methods columns', 'the step is not described in the source text'),
                          ('Population', 'not stated')])
    nts.append(abbrev(['CTEPH', 'GWAS', 'PMID', 'HLA', 'PICS']))
    return spec('4 Study methods', '4  ' + QUESTIONS[3][1],
                'Design › study, with its methods by step. Hover over a methods cell for the source passage.',
                cols, rows, brief=brief, summaries=summary_block(D, 'methods', 'Generated summary of the methods'),
                notes=nts)


# ── 5 condition genetics ─────────────────────────────────────────────────────
def spec_condition_genetics(D):
    cond, lo, men, gc = D['conditions'], D['loci'].copy(), D['mendelian'], D['gc']
    lo['_s'], lo['_p'] = num(lo.l2g1_score), num(lo.log10_p)
    lo = lo[lo._s >= D['l2g_min']].sort_values('_p')
    best = lo.groupby(['node_id', 'l2g1_symbol']).agg(n=('study_locus_id', 'size'), m=('p_mantissa', 'first'),
                                                       e=('p_exponent', 'first'), rs=('rsid', 'first'),
                                                       vid=('variant_id', 'first'), s=('_s', 'first'),
                                                       p=('_p', 'first'), eas=('eas', lambda x: x.map(b).any()))
    best = best.to_dict('index')
    cur = {}
    for x in men.itertuples() if len(men) else []:
        cls = x.classifications or x.clinical_significance
        cur.setdefault((x.node_id, x.symbol), []).append(
            f'{x.source}: {cls}' + (f' ({x.review_status})' if getattr(x, 'review_status', '') else '')
            if cls else f'{x.source}: score {float(x.source_score):.2f}')
    cols = [{'h': 'Category › condition › gene', 'w': 40, 'wrap': True}, {'h': 'GWAS loci', 'w': 7, 'a': 'right',
                                                                         'n': 'a'},
            {'h': 'Best P', 'w': 12, 'a': 'right'}, {'h': 'Lead variant', 'w': 13}, {'h': 'L2G score', 'w': 7,
                                                                                    'a': 'right', 'n': 'b'},
            {'h': 'EAS', 'w': 5, 'a': 'center', 'n': 'c'}, {'h': f'{CURATED} evidence', 'w': 44, 'wrap': True,
                                                             'n': 'd'},
            {'h': 'In CTEPH', 'w': 7, 'a': 'center', 'n': 'e'}, {'h': 'Other related conditions', 'w': 9,
                                                                 'a': 'right', 'n': 'f'}]
    nc = len(cols)
    rows = []
    for cat in CATEGORY_ORDER:
        sub = cond[(cond.category == cat) & cond.node_id.isin(D['in_sheet5'])].sort_values(
            'name', key=lambda s: s.str.lower())
        if not len(sub):
            continue
        rows.append(R(0, 'group0', [f'{cat} ({len(sub)})', CATEGORY_DESC[cat]] + [''] * (nc - 2), span=(1, nc - 1)))
        for c in sub.itertuples():
            g = gc[(gc.node_id == c.node_id)]
            genes = sorted(set(g.symbol))
            shared = [x for x in genes if x in D['cteph_genes']]
            n_g, n_c = g[g.code == 'G'].symbol.nunique(), g[g.code == 'C'].symbol.nunique()
            if genes:
                txt = [f'{plural(int(c.n_gwas_studies), "GWAS study", "GWAS studies")} ({c.n_eas_studies} East '
                       f'Asian), {plural(int(c.n_credible_sets), "credible set")}; {plural(len(genes), "gene")} '
                       f'({n_g} by GWAS, {n_c} curated)'] + \
                    ([', shared with CTEPH: '] + italic_list(shared) if shared else []) + ['.']
            else:
                txt = [f'{plural(int(c.n_gwas_studies), "GWAS study", "GWAS studies")}, '
                       f'{plural(int(c.n_credible_sets), "credible set")}; no gene with a locus-to-gene score ≥ '
                       f'{D["l2g_min"]} or curated evidence.']
            rows.append(R(1, 'group1', [c.name, txt] + [''] * (nc - 2), span=(1, nc - 1),
                          key=('5 Condition genetics', c.node_id)))
            order = sorted(genes, key=lambda x: (x not in D['cteph_genes'], (c.node_id, x) not in cur,
                                                 best[(c.node_id, x)]['p'] if (c.node_id, x) in best else 0, x))
            for x in order:
                bb = best.get((c.node_id, x))
                rel_other = D['rel_n'].get(x, 0) - 1
                cells = [[('i', x)], bb['n'] if bb else '', p_parts(bb['m'], bb['e']) if bb else '',
                         (bb['rs'] or variant_label(bb['vid'])) if bb else '', f'{bb["s"]:.2f}' if bb else '',
                         'yes' if bb and bb['eas'] else '', '\n'.join(cur.get((c.node_id, x), [])) or '',
                         D['cteph_kinds'].get(x, ''), rel_other or '']
                link = ('@', '2 Shared genes', x) if x in D['gene_map_genes'] else None
                rows.append(R(2, 'item', cells, links={0: link} if link else None, hide=True))
    st = D['tidy']
    brief = [[f'{len(D["in_sheet5"])} conditions with genetic data: {st["gwas_studies"]} GWAS studies ({st["eas_studies"]} '
              f'East Asian), {st["credible_sets"]:,} credible sets, {st["mendelian_links"]} curated gene links.'],
             ['A condition (+) opens its genes: those that are the top locus-to-gene gene of one of its credible sets, '
              'and those with curated evidence. Genes also found in CTEPH come first.'],
             ['The full GWAS records are in the appendix (A1 studies, A2 loci, A4 curated genes).']]
    nts = [('a', ['Credible sets of the condition whose top locus-to-gene gene is this gene; P, lead variant and score '
                  'are those of the strongest.']),
           ('b', [f'Locus-to-gene score (0–1), a prediction of the causal gene; genes from {D["l2g_min"]} upward.']),
           ('c', [f'Any of those credible sets from a study with a mostly East Asian sample (≥ {D["eas_min"]:.0%}).']),
           ('d', ['Source: ClinGen or Gene2Phenotype class, PanelApp rating, Orphanet assessment, ClinVar clinical '
                  'significance (review status), UniProt confidence; or the source score when it gives none.']),
           ('e', ['Genetic evidence in CTEPH: G GWAS, L CTEPH literature.']),
           ('f', ['Related conditions other than this one with genetic evidence for the gene.'])]
    nts += [('', [('b', 'Missing values')]), ('', ['   Empty cells: no evidence of that kind.'])]
    nts.append(abbrev(['CTEPH', 'GWAS', 'L2G', 'EAS', 'ClinGen', 'PanelApp']))
    return spec('5 Condition genetics', '5  ' + QUESTIONS[4][1],
                'Category › condition › gene: GWAS loci and curated gene evidence of each related condition.',
                cols, rows, brief=brief, notes=nts)


# ── 6 mechanisms ─────────────────────────────────────────────────────────────
def spec_mechanisms(D):
    m = D['summ'][D['summ'].kind == 'mechanism'].copy()
    m = m.assign(_n=-num(m.n_targets)).sort_values('_n')
    tg = D['targets'].copy()
    tg['lit'] = num(tg.ds_europepmc)
    tg = tg[tg.lit.notna()].sort_values('lit', ascending=False)
    L = D['literature'].assign(_s=lambda d: num(d.evidence_score)).sort_values('_s', ascending=False)
    cols = [{'h': 'Theme › target › publication', 'w': 30, 'wrap': True},
            {'h': 'Count', 'w': 8, 'a': 'right', 'n': 'a'},
            {'h': 'Summary › target › sentences', 'w': 110, 'wrap': True, 'n': 'b'}, {'h': 'Link', 'w': 15}]
    rows = []
    for x in m.itertuples():
        src = set(x.genes.split('; '))
        sub = tg[tg.themes.str.split('; ').map(lambda xs: x.key in xs)]
        txt = md_parts(x.summary) if b(x.grounded) else [f'Summary withheld: {x.problems}.']
        rows.append(R(0, 'group1', [x.title, count(x.n_targets), [('b', 'Generated summary. ')] + txt
                                    if b(x.grounded) else txt, ''], bold=True))
        for t in sub.itertuples():
            ll = L[L.ensembl_id == t.ensembl_id]
            detail = [t.name + f'. Literature score {float(t.lit):.2f}; {len(ll)} publications'
                      + (f' ({year_span(t.first_year, t.last_year)})' if year_span(t.first_year, t.last_year) else '')
                      + '.' + (' A source of the summary.' if t.symbol in src else '')]
            if len(ll) and ll.matched_as.iloc[0]:
                detail += [('g', ' Matched as: ' + ', '.join(sorted({w for v in ll.matched_as for w in v.split(' | ')
                                                                     if w})[:6]) + '.')]
            rows.append(R(1, 'group1', [[('i', t.symbol)], len(ll), detail, ''], hide=True))
            for p in ll.itertuples():
                sents = [s for s in p.sentences.split(' || ') if s]
                body = ([('b', p.title)] if p.title else []) + \
                    [('g', ('\n' if (p.title or k) else '') + '• ' + s) for k, s in enumerate(sents)]
                rows.append(R(2, 'quote', [p.year or None, '', body or [('g', 'only the title names both')],
                                           (f'PMID {p.pmid}' if str(p.pmid).isdigit() else f'Europe PMC {p.pmid}')
                                           if p.pmid else None],
                              links={3: pub_url(p.pmid)} if pub_url(p.pmid) else None, hide=True))
    brief = [[f'{len(tg)} CTEPH targets with text-mined literature, in {len(m)} Reactome themes (a target can be in '
              'several). Each theme opens to its targets and their publications with the sentences that name the '
              'target and CTEPH.'],
             ['The summary of each theme was generated from up to two sentences of the best-scored publications of '
              'its best-scored targets, and checked: every PMID it cites is among those, every gene it names occurs '
              'in them.'],
             ['Text mining can map an abbreviation to the wrong gene (for example TF, tissue factor, mapped to '
              'transferrin): check the words a target was matched as.']]
    nts = [('a', ['Theme rows: targets in the theme; target rows: their publications.']),
           ('b', ['Theme rows: the generated summary; target rows: the target, its literature score and the words it '
                  'was matched as; publication rows: the title and up to three sentences, abstract first.'])]
    nts += missing_block([('the first column', 'the publication has no date'),
                          ('Link', 'the record has no identifier')])
    nts.append(abbrev(['CTEPH', 'PMID']))
    return spec('6 Mechanisms', '6  ' + QUESTIONS[5][1],
                'Reactome theme › text-mined CTEPH target › publication and its sentences.', cols, rows, brief=brief,
                notes=nts)


# ── 7 East Asian ─────────────────────────────────────────────────────────────
def spec_eas(D):
    cols = [{'h': 'Kind › item', 'w': 34, 'wrap': True}, {'h': 'Evidence', 'w': 80, 'wrap': True},
            {'h': 'Population', 'w': 24, 'wrap': True}, {'h': 'Reference', 'w': 16}]
    names = D['names']
    rows = []
    rel = D['relations']
    ea = rel[rel.east_asian.map(b)]
    rows.append(R(0, 'group0', [f'Clinical evidence ({len(ea)})', '', '', ''], span=(0, 3)))
    for x in ea.itertuples():
        rows.append(R(1, 'item', [names.get(x.node_id, x.node_id), [x.detail] + (['\n', ('gi', f'“{x.quote}”')]
                                                                                 if x.quote else []),
                                  x.population or None, x.citation.split(',')[0] if x.pmid else None],
                      links={3: pub_url(x.pmid)} if x.pmid else None))
    cst = D['cst']
    s = cst[cst.eas].sort_values(['year', 'pmid'])
    rows.append(R(0, 'group0', [f'Genetic studies of CTEPH ({len(s)}; Japanese {int(s.jp.sum())})', '', '', ''],
                  span=(0, 3)))
    fx = D['fx']
    for x in s.itertuples():
        f = fx[fx.pmid == x.pmid]
        res = '; '.join(dict.fromkeys(f'{y.label} {y.variant} {y.result}'.replace('  ', ' ') for y in f.itertuples()))
        rows.append(R(1, 'item', [x.cite, [f'{x.design}. ', ('g', res or 'no gene-level finding')],
                                  x.population, f'PMID {x.pmid}'], links={3: pub_url(x.pmid)}))
    st = D['studies']
    e = st[st.eas.map(b)].assign(_n=lambda d: -num(d.n_samples)).sort_values(['condition', '_n'])
    rows.append(R(0, 'group0', [f'GWAS of the related conditions with a mostly East Asian sample ({len(e)})', '',
                                '', ''], span=(0, 3)))
    for x in e.itertuples():
        pub = cite(x) if x.first_author else (x.project or None)
        rows.append(R(1, 'item', [x.condition, f'{x.trait}; {intx(x.n_samples) or "sample size not given"} samples; '
                                  f'{x.n_credible_sets} credible sets ({x.study_id}).', ancestry_label(x.ancestry)
                                  or None, pub], links={3: pub_url(x.pmid)} if pub_url(x.pmid) else None))
    co = D['coloc']
    c = co[co.eas.map(b)].assign(_c=lambda d: -num(d.clpp)).sort_values(['cteph_locus', '_c'])
    rows.append(R(0, 'group0', [f'GWAS signals from East Asian studies that colocalise with a CTEPH locus ({len(c)})',
                                '', '', ''], span=(0, 3)))
    for x in c.itertuples():
        rows.append(R(1, 'item', [x.condition or x.trait, [f'Colocalises with the CTEPH locus {x.cteph_locus} (',
                                                         ('i', x.cteph_gene), f'), CLPP {float(x.clpp):.2f}; '
                                                         f'{x.trait} ({x.study_id}).'],
                                  ancestry_label(x.ancestry) or None,
                                  ' '.join(y for y in (x.first_author, x.pub_date[:4]) if y) or x.project or None],
                      links={3: pub_url(x.pmid)} if pub_url(x.pmid) else None))
    brief = [['Evidence from East Asian populations, gathered from the other sheets: clinical evidence quoted from '
              'PubMed, genetic studies of CTEPH done in Japan or another East Asian country (or so flagged by the '
              f'extraction), and GWAS whose sample is at least {D["eas_min"]:.0%} East Asian.']]
    nts = missing_block([('Population', 'not stated'), ('Reference', 'no publication listed (consortium release)')])
    nts.append(abbrev(['CTEPH', 'GWAS', 'PMID', 'CLPP']))
    return spec('7 East Asian', '7  ' + QUESTIONS[6][1], 'Kind of evidence › item.', cols, rows, brief=brief,
                notes=nts)


# ── 8 changes ────────────────────────────────────────────────────────────────
def s_changes(wb, D, foot):
    sh = new_sheet(wb, '8 Changes', '#1F4E79')
    c, ch = D['changes'], D['changes_json']
    sub = (f'Compared with the snapshot of {ch["previous"]} (release {ch["previous_release"]})' if ch['previous']
           else 'First snapshot: nothing to compare with')
    title(sh, '8  ' + QUESTIONS[7][1], sub)
    sh.ws.write_url(0, 3, "internal:'Guide'!A1", sh.fmt(font_color=LINK, underline=1, align='right'),
                    string='← Guide')
    cols = [{'h': 'Category', 'w': 20}, {'h': 'Change', 'w': 10}, {'h': 'Item', 'w': 44, 'wrap': True},
            {'h': 'Detail', 'w': 80, 'wrap': True}]
    rows = [[x.category, x.change, x.item, x.detail or None] for x in c.itertuples()]
    r = table(sh, 3, cols, rows, freeze=0, autofilter=True)
    items = missing_block([('Detail', 'no further detail recorded')]) if len(c) and c.detail.eq('').any() else []
    r = notes(sh, r + 1, items or [('', [' '])])
    print_setup(sh.ws, r, len(cols), foot)
    return {'header': 3, 'rows': len(c), 'spanner': False}


# ── appendix: flat tables with filters ───────────────────────────────────────
def flat(wb, name, sub, cols, rows, links, nts, foot, freeze=1):
    sh = new_sheet(wb, name, '#A6A6A6')
    title(sh, dict(APPENDIX)[name], sub)
    sh.ws.write_url(0, len(cols) - 1, "internal:'Guide'!A1", sh.fmt(font_color=LINK, underline=1, align='right'),
                    string='← Guide')
    r = table(sh, 3, cols, rows, links, freeze=freeze, autofilter=True)
    r = notes(sh, r + 1, nts)
    print_setup(sh.ws, r, len(cols), foot)
    return {'header': 3, 'rows': len(rows), 'spanner': False}


def a_studies(wb, D, foot):
    s = D['studies'].copy()
    s['_c'] = (s.node_id != CORE).astype(int)
    s['_n'] = num(s.n_samples)
    s = s.sort_values(['_c', 'condition', '_n'], ascending=[True, True, False])
    cols = [{'h': 'Condition', 'w': 28, 'wrap': True}, {'h': 'Study', 'w': 22}, {'h': 'Trait', 'w': 30, 'wrap': True},
            {'h': 'Publication', 'w': 18}, {'h': 'PMID', 'w': 9}, {'h': 'Samples', 'w': 10, 'a': 'right'},
            {'h': 'Cases', 'w': 9, 'a': 'right'}, {'h': 'Controls', 'w': 10, 'a': 'right'},
            {'h': 'Cohorts', 'w': 22, 'wrap': True}, {'h': 'Ancestry', 'w': 22, 'wrap': True, 'n': 'a'},
            {'h': 'EAS', 'w': 6, 'a': 'center', 'n': 'b'}, {'h': 'Summary statistics', 'w': 10, 'a': 'center'},
            {'h': 'Credible sets', 'w': 9, 'a': 'right'}]
    rows, links = [], {}
    for i, x in enumerate(s.itertuples()):
        pub = ' '.join(y for y in (x.first_author, x.year) if y) or x.project or None
        rows.append([x.condition, x.study_id, x.trait or None, pub, x.pmid or None, intx(x.n_samples),
                     intx(x.n_cases), intx(x.n_controls), x.cohorts or None, ancestry_label(x.ancestry) or None,
                     'yes' if b(x.eas) else 'no', 'yes' if b(x.summary_statistics) else 'no', x.n_credible_sets])
        if pub_url(x.pmid):
            links[(i, 4)] = pub_url(x.pmid)
    nts = [('a', ['Population of the study\'s linkage-disequilibrium reference, with its share of the sample.']),
           ('b', [f'East Asian share of the sample ≥ {D["eas_min"]:.0%}.'])]
    nts += missing_block([('Trait, Publication, PMID, Samples, Cases, Controls, Cohorts, Ancestry',
                           'not given by Open Targets for the study')])
    nts.append(abbrev(['CTEPH', 'GWAS', 'EAS', 'PMID']))
    return flat(wb, 'A1 GWAS studies', f'{len(s)} studies, CTEPH first; use the filter buttons of the header row',
                cols, rows, links, nts, foot, freeze=2)


def a_loci(wb, D, foot):
    g = D['loci'].copy()
    g['_c'] = (g.node_id != CORE).astype(int)
    g['_p'] = num(g.log10_p)
    g = g.sort_values(['_c', 'condition', '_p'])
    cols = [{'h': 'Condition', 'w': 26, 'wrap': True}, {'h': 'Study', 'w': 20}, {'h': 'Publication', 'w': 16},
            {'h': 'Lead variant', 'w': 13}, {'h': 'Position (GRCh38)', 'w': 22},
            {'h': 'P', 'w': 12, 'a': 'right', 'n': 'a'}, {'h': 'β', 'w': 8, 'a': 'right', 'n': 'a'},
            {'h': 'Fine-mapping', 'w': 10}, {'h': 'EAS', 'w': 6, 'a': 'center'},
            {'h': 'Locus-to-gene genes (score)', 'w': 30, 'wrap': True, 'n': 'b'}]
    rows = []
    for x in g.itertuples():
        l2g = [(getattr(x, f'l2g{k}_symbol'), num(pd.Series([getattr(x, f'l2g{k}_score')]))[0]) for k in (1, 2, 3)
               if getattr(x, f'l2g{k}_symbol')]
        be = num(pd.Series([x.beta]))[0]
        pub = ' '.join(y for y in (x.first_author, x.year) if y) or x.project or None
        rows.append([x.condition, x.study_id, pub, x.rsid or None, variant_label(x.variant_id) or None,
                     p_parts(x.p_mantissa, x.p_exponent), None if pd.isna(be) else f'{be:.3f}', x.finemapping or None,
                     'yes' if b(x.eas) else 'no', genes_parts(l2g) or None])
    nts = [('a', ['P and β of the credible set\'s lead variant, as reported by the study.']),
           ('b', ['Up to three genes by Open Targets locus-to-gene score (0–1), a prediction of the causal gene.'])]
    nts += missing_block([('Publication', 'none listed (consortium release)'), ('Lead variant', 'no rsID'),
                          ('β, Fine-mapping', 'not given'), ('Locus-to-gene genes', 'no prediction')])
    nts.append(abbrev(['CTEPH', 'GWAS', 'L2G', 'EAS', 'β']))
    return flat(wb, 'A2 GWAS loci', f'{len(g):,} credible sets, by condition and P', cols, rows, {}, nts, foot)


def a_coloc(wb, D, foot):
    c = D['coloc'].copy()
    c['_p'] = -num(c.clpp)
    c = c.sort_values(['cteph_locus', '_p'])
    cols = [{'h': 'CTEPH locus', 'w': 12}, {'h': 'CTEPH gene', 'w': 9, 'n': 'a'},
            {'h': 'Condition / trait', 'w': 30, 'wrap': True}, {'h': 'Study', 'w': 20},
            {'h': 'Trait as reported', 'w': 30, 'wrap': True}, {'h': 'Publication', 'w': 16},
            {'h': 'Ancestry', 'w': 18, 'wrap': True}, {'h': 'EAS', 'w': 6, 'a': 'center'},
            {'h': 'Lead variant', 'w': 13}, {'h': 'P', 'w': 12, 'a': 'right'},
            {'h': 'CLPP', 'w': 7, 'a': 'right', 'n': 'b'}, {'h': 'Method', 'w': 16}]
    rows = []
    for x in c.itertuples():
        pub = ' '.join(y for y in (x.first_author, x.pub_date[:4]) if y) or x.project or None
        rows.append([x.cteph_locus, [('i', x.cteph_gene)] if x.cteph_gene else None, x.condition or None, x.study_id,
                     x.trait or None, pub, ancestry_label(x.ancestry) or None, 'yes' if b(x.eas) else 'no',
                     x.other_lead or None, p_parts(x.p_mantissa, x.p_exponent),
                     f'{float(x.clpp):.3f}' if x.clpp else None,
                     x.method.replace('COLOC_PIP_ECAVIAR', 'eCAVIAR').replace('COLOC', 'coloc')])
    nts = [('a', ['Top locus-to-gene gene of the CTEPH credible set.']),
           ('b', ['eCAVIAR colocalisation posterior probability of the two fine-mapped signals.'])]
    nts += missing_block([('Condition / trait', 'the study has no disease mapping'),
                          ('Publication, Ancestry, Lead variant', 'not given')])
    nts.append(abbrev(['CTEPH', 'GWAS', 'CLPP', 'EAS']))
    return flat(wb, 'A3 Colocalisation', f'{len(c)} colocalisations', cols, rows, {}, nts, foot)


def a_curated(wb, D, foot):
    m = D['mendelian'].sort_values(['condition', 'symbol', 'source']) if len(D['mendelian']) else D['mendelian']
    cols = [{'h': 'Condition', 'w': 28, 'wrap': True}, {'h': 'Gene', 'w': 10}, {'h': 'Source', 'w': 16},
            {'h': 'Score', 'w': 7, 'a': 'right', 'n': 'a'}, {'h': 'Classification', 'w': 16, 'wrap': True, 'n': 'b'},
            {'h': 'ClinVar clinical significance', 'w': 24, 'wrap': True},
            {'h': 'ClinVar review status', 'w': 26, 'wrap': True}, {'h': 'Specific diseases', 'w': 40, 'wrap': True},
            {'h': 'Records', 'w': 8, 'a': 'right'}, {'h': 'PMIDs', 'w': 24, 'wrap': True}]
    rows = []
    for x in m.itertuples():
        cv = x.source == 'ClinVar'
        rows.append([x.condition, [('i', x.symbol)], x.source, fmt2(x.source_score),
                     '' if cv else (x.classifications or None), (x.clinical_significance or None) if cv else '',
                     (getattr(x, 'review_status', '') or None) if cv else '', x.specific_diseases or None,
                     x.n_records, ' '.join(x.pmids.split()[:6]) + (' …' if len(x.pmids.split()) > 6 else '') or None])
    nts = [('a', ['Open Targets score of the source for the condition and its descendants (0–1); genes from 0.5 '
                  'upward are listed.']),
           ('b', ['ClinGen or Gene2Phenotype class, PanelApp rating (green, amber), Orphanet assessment, UniProt '
                  'confidence, as given.'])]
    nts += [('', [('b', 'Missing values')]),
            ('', ['   Empty cells: the field belongs to another source (ClinVar fields for ClinVar rows only).'])]
    nts += missing_block([('Classification', 'the source gives none (gene burden)'),
                          ('Specific diseases', 'none given'), ('PMIDs', 'none given')])
    nts.append(abbrev(['ClinGen', 'PanelApp', 'PMID']))
    return flat(wb, 'A4 Curated genes', f'{len(m)} condition × gene × source', cols, rows, {}, nts, foot, freeze=2)


def a_cteph_studies(wb, D, foot):
    s = D['cst'].assign(_y=lambda d: num(d.year)).sort_values(['_y', 'pmid'])
    cols = [{'h': 'PMID', 'w': 9}, {'h': 'Year', 'w': 6}, {'h': 'First author', 'w': 14},
            {'h': 'Title', 'w': 44, 'wrap': True}, {'h': 'Study type', 'w': 12}, {'h': 'Design', 'w': 22, 'wrap': True},
            {'h': 'Population', 'w': 12, 'wrap': True}, {'h': 'Japanese', 'w': 8, 'a': 'center', 'n': 'a'},
            {'h': 'East Asian', 'w': 8, 'a': 'center'}, {'h': 'Cases', 'w': 7, 'a': 'right'},
            {'h': 'Controls', 'w': 8, 'a': 'right'}, {'h': 'Control type', 'w': 20, 'wrap': True},
            {'h': 'Genes', 'w': 22, 'wrap': True, 'n': 'b'}, {'h': 'Variants', 'w': 20, 'wrap': True},
            {'h': 'Replication', 'w': 9}, {'h': 'Methods from', 'w': 9, 'a': 'center'}]
    rows, links = [], {}
    for i, x in enumerate(s.itertuples()):
        rows.append([x.pmid, x.year, x.first_author, x.title, x.study_type or None, x.design, x.population,
                     'yes' if x.jp else 'no', 'yes' if x.eas else 'no', x.n_cases or '', x.n_controls or '',
                     x.control_type or '', italic_list(x.gene_symbols.split('; ')) if x.gene_symbols else
                     (x.genes or ''), x.variants or '', x.replication, x.methods_source or None])
        links[(i, 0)] = pub_url(x.pmid)
    nts = [('a', ['Flagged by the extraction, or the study was done in Japan.']),
           ('b', ['Official gene symbols checked against Open Targets; the record\'s own names where none was '
                  'found.'])]
    nts += [('', [('b', 'Missing values')]),
            ('', ['   Empty cells: not stated in the abstract, or no gene named.'])]
    nts += missing_block([('Study type', 'not classified')])
    nts.append(abbrev(['CTEPH', 'PMID', 'HLA']))
    return flat(wb, 'A5 CTEPH studies', f'{len(s)} studies, by year', cols, rows, links, nts, foot)


def a_screening(wb, D, foot):
    s = D['csc'].assign(_y=lambda d: num(d.year), _i=lambda d: d.include != 'True').sort_values(['_i', '_y', 'pmid'])
    cols = [{'h': 'PMID', 'w': 9}, {'h': 'Year', 'w': 6}, {'h': 'Title', 'w': 70, 'wrap': True},
            {'h': 'Included', 'w': 8, 'a': 'center'}, {'h': 'Record type', 'w': 16},
            {'h': 'Decided by', 'w': 12, 'n': 'a'}, {'h': 'Reason', 'w': 60, 'wrap': True}]
    rows, links = [], {}
    for i, x in enumerate(s.itertuples()):
        rows.append([x.pmid, x.year or None, x.title or None, 'yes' if x.include == 'True' else 'no',
                     x.study_type or '', {'rule': 'pre-screen', 'model': 'language model'}.get(x.stage, x.stage),
                     x.reason.replace('pre-screen: ', '') or None])
        links[(i, 0)] = pub_url(x.pmid)
    nts = [('a', ['Pre-screen: a deterministic rule (publication type, no genetic-study term in the title or '
                  'abstract); language model: the remaining records, read in full abstract.'])]
    nts += [('', [('b', 'Missing values')]), ('', ['   Empty cells: the pre-screen does not classify the record.'])]
    nts += missing_block([('Year, Title', 'not in the PubMed record'), ('Reason', 'none recorded')])
    nts.append(abbrev(['CTEPH', 'PMID']))
    return flat(wb, 'A6 Screening', f'{len(s)} PubMed records: included first', cols, rows, links, nts, foot)


# ── checks ───────────────────────────────────────────────────────────────────
def verify(path, stats):
    """Re-open the workbook: sheets; body rows; header letters against footnotes; a reason for every '—'; no
    missing-value artefacts or internal names, in cells or notes; outline levels; internal links; the XML Excel is
    strict about."""
    from openpyxl import load_workbook
    with warnings.catch_warnings():
        warnings.simplefilter('ignore', UserWarning)
        wb = load_workbook(path, rich_text=True)
    assert wb.sheetnames == SH, wb.sheetnames
    for sheet in wb.worksheets:
        L = stats['layout'][sheet.title]
        h = L['header'] + 1
        head = {ch for c in range(1, sheet.max_column + 1) for ch in plain(sheet.cell(h, c).value, sup=True)}
        n = max(L['rows'], 1)
        body = range(h + 1, h + 1 + n)
        assert all(sheet.cell(r, 1).value is not None for r in body), (sheet.title, 'a body row without a label')
        assert sheet.cell(h + 1 + n, 1).value is None, (sheet.title, 'more body rows than recorded')
        letters, notes_text, has_dash = set(), '', False
        for r in range(h + 1 + n, sheet.max_row + 1):
            v = sheet.cell(r, 1).value
            if v is None:
                continue
            sup, txt = plain(v, sup=True), plain(v)
            if sup and txt.startswith(' ') and len(sup) == 1:
                letters.add(sup)
            notes_text += txt
        for row in sheet.iter_rows(min_row=h + 1, max_row=h + n):
            for c in row:
                if c.value is not None and plain(c.value) == MISSING:
                    has_dash = True
        assert head <= letters, (sheet.title, sorted(head), sorted(letters))
        if has_dash:
            assert '—,' in notes_text or '—:' in notes_text, (sheet.title, 'a — without its reason')
        prev = 0
        for r in body:
            d = sheet.row_dimensions[r]
            lvl = d.outlineLevel or 0
            assert lvl <= prev + 1, (sheet.title, r, 'outline level jumps')
            assert not d.hidden or lvl >= 1, (sheet.title, r, 'a hidden row outside a group')
            prev = lvl
        for row in sheet.iter_rows():
            for c in row:
                texts = [] if c.value is None else [plain(c.value) + plain(c.value, sup=True)]
                if c.comment is not None:
                    texts.append(c.comment.text)
                for s in texts:
                    assert s not in ('nan', 'None', 'NaN') and 'nan' not in s.split() and 'NaN' not in s.split(), \
                        (sheet.title, c.coordinate, s[:80])
                    bad = [t for t in FORBIDDEN_TOKENS if t in s]
                    assert not bad, (sheet.title, c.coordinate, bad, s[:80])
                if c.hyperlink is not None and c.hyperlink.location:
                    m = re.match(r"'?([^'!]+)'?!A(\d+)$", c.hyperlink.location)
                    assert m and m.group(1) in SH, (sheet.title, c.coordinate, c.hyperlink.location)
                    tgt = wb[m.group(1)]
                    assert tgt.cell(int(m.group(2)), 1).value is not None, (sheet.title, c.coordinate,
                                                                          c.hyperlink.location, 'empty target')
    z = zipfile.ZipFile(path)
    for name in z.namelist():
        if name.startswith('xl/worksheets/sheet'):
            x = z.read(name).decode()
            assert 'inlineStr' not in x, f'inline strings in {name}'
            for pane in re.findall(r'<pane [^>]*>', x):
                assert 'xSplit=' in pane or 'ySplit=' in pane, (name, pane)
    ss = z.read('xl/sharedStrings.xml').decode()
    bare = [t for p, t in re.findall(r'<t( xml:space="preserve")?>([^<]*)</t>', ss) if not p and t != t.strip()]
    assert not bare, bare[:3]


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--tables', required=True)
    ap.add_argument('--summaries', required=True)
    ap.add_argument('--changes', required=True)
    ap.add_argument('--changes-json', required=True)
    ap.add_argument('--knowledge-json', required=True, help='for the caveats')
    ap.add_argument('--themes', required=True, help='gene themes of the CTEPH literature (YAML)')
    ap.add_argument('--l2g-min', type=float, default=0.5)
    ap.add_argument('--eas-min', type=float, default=0.5)
    ap.add_argument('--out', required=True)
    ap.add_argument('--out-stats', default='workbook.json')
    a = ap.parse_args()
    T = Path(a.tables)
    cond = read_tsv(T / 'conditions.tsv')
    D = {'tidy': json.loads((T / 'tidy.json').read_text()), 'conditions': cond, 'relations': read_tsv(T / 'relations.tsv'),
         'coloc': read_tsv(T / 'colocalisation.tsv'), 'studies': read_tsv(T / 'gwas_studies.tsv'),
         'loci': read_tsv(T / 'gwas_loci.tsv'), 'mendelian': read_tsv(T / 'mendelian.tsv'),
         'gene_condition': read_tsv(T / 'gene_condition.tsv'), 'literature': read_tsv(T / 'literature.tsv'),
         'targets': read_tsv(T / 'cteph_targets.tsv'), 'cst': read_tsv(T / 'cteph_studies.tsv'),
         'cfi': read_tsv(T / 'cteph_findings.tsv'), 'cmf': read_tsv(T / 'cteph_methods.tsv'),
         'csc': read_tsv(T / 'cteph_screening.tsv'), 'summ': read_tsv(a.summaries),
         'changes': read_tsv(a.changes), 'changes_json': json.loads(Path(a.changes_json).read_text()),
         'caveats': json.loads(Path(a.knowledge_json).read_text())['caveats'], 'themes': load_themes(a.themes),
         'names': dict(zip(cond.node_id, cond.name)), 'l2g_min': a.l2g_min, 'eas_min': a.eas_min}
    prepare(D)
    st = D['tidy']
    foot = f'CTEPH knowledge graph · Open Targets {st["release"]} · fetched {st["fetch_date"]}'
    # every question sheet is laid out before any is written, so that sheets can link to each other's rows
    D['gene_map_genes'] = set()
    gmap = spec_genes(D)
    D['gene_map_genes'] = {r['key'][1] for r in gmap['rows'] if r['key']}
    specs = [spec_conditions(D), gmap, spec_cteph(D), spec_methods(D), spec_condition_genetics(D),
             spec_mechanisms(D), spec_eas(D)]
    anchors = {}
    for s in specs:
        top, keys = layout(s)
        assert not set(keys) & set(anchors), set(keys) & set(anchors)
        anchors.update(keys)
    D['appendix_rows'] = {'A1 GWAS studies': len(D['studies']), 'A2 GWAS loci': len(D['loci']),
                          'A3 Colocalisation': len(D['coloc']), 'A4 Curated genes': len(D['mendelian']),
                          'A5 CTEPH studies': len(D['cst']), 'A6 Screening': len(D['csc'])}
    wb = xlsxwriter.Workbook(a.out, {'nan_inf_to_errors': False})
    wb.set_properties({'title': f'CTEPH knowledge graph ({st["fetch_date"]})',
                       'subject': 'Conditions related to CTEPH and their genetics'})
    lay = {'Guide': s_guide(wb, D, foot)}
    for s in specs:
        lay[s['name']] = write_spec(wb, s, anchors, foot)
    lay['8 Changes'] = s_changes(wb, D, foot)
    for (name, _d), f in zip(APPENDIX, (a_studies, a_loci, a_coloc, a_curated, a_cteph_studies, a_screening)):
        lay[name] = f(wb, D, foot)
    wb.close()
    stats = {'sheets': SH, 'rows': {k: v['rows'] for k, v in lay.items()}, 'layout': lay, 'release': st['release'],
             'fetch_date': st['fetch_date']}
    verify(a.out, stats)
    Path(a.out_stats).write_text(json.dumps(stats, indent=1))
    log(f'workbook: {", ".join(f"{k} {v}" for k, v in stats["rows"].items())}')


if __name__ == '__main__':
    main()
